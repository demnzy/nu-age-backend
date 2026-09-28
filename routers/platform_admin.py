import os
import sys
import io
import csv
import time
from uuid import UUID
from datetime import datetime, timezone, timedelta
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from fastapi import APIRouter, Depends, HTTPException, status, Query, BackgroundTasks
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, EmailStr
from sqlalchemy import or_, and_, func, desc, String, cast
from sqlalchemy.orm import Session

from database import get_db, Settings
import models
from schemas import Roles
from services import auth, utils
from services.notifications import send_push_notification
import resend


router = APIRouter(prefix="/platform-admin", tags=["Platform Super Admin"])


# ─────────────────────────────────────────────────────────────────────────────
# PYDANTIC SCHEMAS
# ─────────────────────────────────────────────────────────────────────────────

class AdminLoginRequest(BaseModel):
    username: str  # Can be username or email
    password: str


class AdminVerifyResponse(BaseModel):
    access_token: str
    token_type: str = "Bearer"
    user_id: UUID
    username: str
    email: str
    role: str
    name: str


class PlatformAnalyticsResponse(BaseModel):
    total_users: int
    total_students: int
    total_teachers: int
    total_admins: int
    verified_users: int
    unverified_users: int
    active_today: int
    active_this_week: int
    total_organisations: int
    total_courses: int
    total_enrollments: int
    total_device_tokens: int


class AdminUserItem(BaseModel):
    id: UUID
    name: str
    first_name: str
    last_name: str
    username: str
    email: str
    number: Optional[str] = None
    role: str
    gender: str
    university: Optional[str] = None
    streak: int = 0
    is_verified: bool = False
    last_login_date: Optional[str] = None
    created_at: Optional[str] = None
    enrolled_courses_count: int = 0
    created_courses_count: int = 0
    organisations_count: int = 0


class AdminUsersPageResponse(BaseModel):
    items: List[AdminUserItem]
    total: int
    page: int
    limit: int
    total_pages: int


class AdminUserUpdateRequest(BaseModel):
    role: Optional[str] = None
    is_verified: Optional[bool] = None
    streak: Optional[int] = None
    university: Optional[str] = None


class BulkEmailRequest(BaseModel):
    audience: str = "all"  # "all", "students", "teachers", "admins", "unverified"
    subject: str
    body_html: str
    sender_name: Optional[str] = "Tobi from Nu Age"


class BulkPushRequest(BaseModel):
    audience: str = "all"  # "all", "students", "teachers", "admins", "unverified", "test_me"
    title: str
    body: str
    subtitle: Optional[str] = None
    image_url: Optional[str] = None
    action_route: Optional[str] = None
    action_button_label: Optional[str] = None
    priority: int = 10  # 10=high, 5=normal
    ttl_seconds: int = 86400  # 24h default


class BroadcastHistoryItem(BaseModel):
    id: int
    sender_username: Optional[str] = None
    title: str
    body: str
    subtitle: Optional[str] = None
    image_url: Optional[str] = None
    action_route: Optional[str] = None
    action_button_label: Optional[str] = None
    audience: str
    targeted_devices_count: int = 0
    delivered_count: int = 0
    failed_count: int = 0
    priority: int = 10
    ttl_seconds: int = 86400
    status: str
    created_at: Optional[str] = None


class BroadcastHistoryResponse(BaseModel):
    items: List[BroadcastHistoryItem]
    total: int
    page: int
    limit: int



# ─────────────────────────────────────────────────────────────────────────────
# SUPER ADMIN DEPENDENCY GUARD
# ─────────────────────────────────────────────────────────────────────────────

def get_configured_platform_admins() -> List[str]:
    """
    Retrieves the list of designated super admin usernames and emails
    using the codebase Settings() env search pattern in database.py
    with os.getenv fallback.
    """
    try:
        raw = getattr(Settings(), "PLATFORM_SUPER_ADMINS", "")
    except Exception:
        raw = ""
    if not raw:
        raw = os.getenv("PLATFORM_SUPER_ADMINS", "")
    admins = [s.strip().lower() for s in str(raw).split(",") if s.strip()]
    if "nu-admin" not in admins:
        admins.append("nu-admin")
    return admins


def get_current_super_admin(
    current_user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
) -> models.User:
    """
    Enforces that the caller is authenticated and possesses Platform Super Admin
    privileges (either role == Admin, or listed in PLATFORM_SUPER_ADMINS).
    """
    role_str = str(getattr(current_user, "role", "")).upper()
    is_admin_role = (getattr(current_user, "role", None) == Roles.ADMIN or "ADMIN" in role_str)

    configured_admins = get_configured_platform_admins()
    is_designated = (
        current_user.username.lower() in configured_admins
        or (current_user.email and current_user.email.lower() in configured_admins)
    )

    if not (is_admin_role or is_designated):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Access forbidden: Platform Super Administrator privileges required.",
        )

    return current_user


# ─────────────────────────────────────────────────────────────────────────────
# 1. AUTHENTICATION & LOCK-SCREEN VERIFICATION
# ─────────────────────────────────────────────────────────────────────────────

@router.post("/auth/verify", response_model=AdminVerifyResponse)
def verify_super_admin_credentials(
    payload: AdminLoginRequest,
    db: Session = Depends(get_db),
):
    """
    Verifies super-admin credentials directly from the admin lock screen.
    Returns an elevated access token and admin profile.
    """
    ident = payload.username.strip()
    user = (
        db.query(models.User).filter(models.User.email == ident).first()
        or db.query(models.User).filter(models.User.username == ident).first()
    )

    if not user:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User account not found.")

    if not utils.verify_password(payload.password, user.password):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Incorrect password.")

    role_str = str(getattr(user, "role", "")).upper()
    is_admin_role = (getattr(user, "role", None) == Roles.ADMIN or "ADMIN" in role_str)
    configured_admins = get_configured_platform_admins()
    is_designated = (
        user.username.lower() in configured_admins
        or (user.email and user.email.lower() in configured_admins)
    )

    if not (is_admin_role or is_designated):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="This account is not designated as a Nu-Age Platform Administrator.",
        )

    # Issue access token with is_super_admin payload
    token = auth.create_access_token({"email": user.email, "is_super_admin": True})

    full_name = f"{user.first_name} {user.last_name}".strip()
    return {
        "access_token": token["access_token"],
        "token_type": "Bearer",
        "user_id": user.id,
        "username": user.username,
        "email": user.email,
        "role": str(user.role.value if hasattr(user.role, "value") else user.role),
        "name": full_name or user.username,
    }


# ─────────────────────────────────────────────────────────────────────────────
# 2. PLATFORM ANALYTICS & TELEMETRY
# ─────────────────────────────────────────────────────────────────────────────

@router.get("/analytics", response_model=PlatformAnalyticsResponse)
def get_platform_analytics(
    current_admin: models.User = Depends(get_current_super_admin),
    db: Session = Depends(get_db),
):
    """
    Computes global platform KPIs: total users, roles breakdown, verification status,
    active users, total organisations, courses, and enrollments.
    """
    now = datetime.now(timezone.utc)
    today = now.date()
    seven_days_ago = today - timedelta(days=7)

    total_users = db.query(func.count(models.User.id)).scalar() or 0
    total_students = db.query(func.count(models.User.id)).filter(models.User.role == Roles.STUDENT).scalar() or 0
    total_teachers = db.query(func.count(models.User.id)).filter(models.User.role == Roles.TEACHER).scalar() or 0
    total_admins = db.query(func.count(models.User.id)).filter(models.User.role == Roles.ADMIN).scalar() or 0

    verified_users = db.query(func.count(models.User.id)).filter(models.User.is_verified == True).scalar() or 0
    unverified_users = total_users - verified_users

    active_today = db.query(func.count(models.User.id)).filter(models.User.last_login_date == today).scalar() or 0
    active_this_week = db.query(func.count(models.User.id)).filter(models.User.last_login_date >= seven_days_ago).scalar() or 0

    total_orgs = db.query(func.count(models.Organisation.id)).scalar() or 0
    total_courses = db.query(func.count(models.Course.id)).scalar() or 0

    # Total enrollments
    enrollment_table = models.Base.metadata.tables.get("enrollments")
    total_enrollments = 0
    if enrollment_table is not None:
        total_enrollments = db.query(func.count()).select_from(enrollment_table).scalar() or 0

    total_device_tokens = db.query(func.count(models.DeviceToken.id)).scalar() or 0

    return {
        "total_users": total_users,
        "total_students": total_students,
        "total_teachers": total_teachers,
        "total_admins": total_admins,
        "verified_users": verified_users,
        "unverified_users": unverified_users,
        "active_today": active_today,
        "active_this_week": active_this_week,
        "total_organisations": total_orgs,
        "total_courses": total_courses,
        "total_enrollments": total_enrollments,
        "total_device_tokens": total_device_tokens,
    }


# ─────────────────────────────────────────────────────────────────────────────
# 3. USER DIRECTORY & MANAGEMENT
# ─────────────────────────────────────────────────────────────────────────────

@router.get("/users", response_model=AdminUsersPageResponse)
def list_platform_users(
    q: Optional[str] = Query(None, description="Search query across name, username, email, university"),
    role: Optional[str] = Query(None, description="Filter by role: Student, Teacher, Admin"),
    is_verified: Optional[bool] = Query(None, description="Filter by verification state"),
    page: int = Query(1, ge=1),
    limit: int = Query(25, ge=1, le=100),
    sort_by: str = Query("created_desc", description="created_desc, created_asc, streak_desc, name_asc"),
    current_admin: models.User = Depends(get_current_super_admin),
    db: Session = Depends(get_db),
):
    """
    Returns a paginated list of users with search and filter controls.
    """
    query = db.query(models.User)

    if isinstance(q, str) and q.strip():
        term = f"%{q.strip().lower()}%"
        query = query.filter(
            or_(
                func.lower(models.User.first_name).like(term),
                func.lower(models.User.last_name).like(term),
                func.lower(func.concat(func.coalesce(models.User.first_name, ''), ' ', func.coalesce(models.User.last_name, ''))).like(term),
                func.lower(models.User.username).like(term),
                func.lower(models.User.email).like(term),
                func.lower(models.User.university).like(term),
            )
        )

    if isinstance(role, str) and role.strip() and role.lower() != "all":
        r_clean = role.strip().lower()
        matched_enum = None
        for r_enum in Roles:
            if r_enum.value.lower() == r_clean or r_enum.name.lower() == r_clean:
                matched_enum = r_enum
                break
        if matched_enum:
            query = query.filter(
                or_(
                    models.User.role == matched_enum,
                    models.User.role == matched_enum.value,
                    func.lower(func.cast(models.User.role, String)) == r_clean,
                )
            )
        else:
            query = query.filter(func.lower(func.cast(models.User.role, String)) == r_clean)

    if is_verified is not None:
        if is_verified is True:
            query = query.filter(models.User.is_verified == True)
        else:
            query = query.filter(or_(models.User.is_verified == False, models.User.is_verified.is_(None)))

    if sort_by == "created_asc":
        query = query.order_by(models.User.created_at.asc())
    elif sort_by == "streak_desc":
        query = query.order_by(models.User.streak.desc())
    elif sort_by == "name_asc":
        query = query.order_by(models.User.first_name.asc(), models.User.last_name.asc())
    else:
        query = query.order_by(models.User.created_at.desc())

    total = query.count()
    total_pages = max(1, (total + limit - 1) // limit)
    offset = (page - 1) * limit
    users = query.offset(offset).limit(limit).all()

    items = []
    for u in users:
        role_val = str(u.role.value if hasattr(u.role, "value") else u.role)
        gender_val = str(u.gender.value if hasattr(u.gender, "value") else u.gender)
        full_name = f"{u.first_name} {u.last_name}".strip()

        items.append(
            AdminUserItem(
                id=u.id,
                name=full_name or u.username,
                first_name=u.first_name,
                last_name=u.last_name,
                username=u.username,
                email=u.email,
                number=u.number,
                role=role_val,
                gender=gender_val,
                university=u.university,
                streak=u.streak or 0,
                is_verified=bool(u.is_verified),
                last_login_date=u.last_login_date.strftime("%Y-%m-%d") if u.last_login_date else None,
                created_at=u.created_at.strftime("%Y-%m-%d %H:%M") if u.created_at else None,
                enrolled_courses_count=len(u.courses) if hasattr(u, "courses") and u.courses else 0,
                created_courses_count=len(u.created_courses) if hasattr(u, "created_courses") and u.created_courses else 0,
                organisations_count=len(u.organisations) if hasattr(u, "organisations") and u.organisations else 0,
            )
        )

    return {
        "items": items,
        "total": total,
        "page": page,
        "limit": limit,
        "total_pages": total_pages,
    }


# ─────────────────────────────────────────────────────────────────────────────
# 3B. EXCEL (.XLSX) & CSV DATA EXPORT (Declared before {user_id} to prevent 422)
# ─────────────────────────────────────────────────────────────────────────────

def _generate_excel_workbook(users: List[models.User]) -> io.BytesIO:
    """
    Generates a beautifully styled Excel workbook containing all user information.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nu-Age Users"

    # Brand Colors
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # Slate 800
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="0F172A")
    border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    headers = [
        "User ID",
        "Full Name",
        "First Name",
        "Last Name",
        "Username",
        "Email",
        "Phone Number",
        "Role",
        "Gender",
        "Verified",
        "Streak (Days)",
        "University",
        "Last Login Date",
        "Created At (UTC)",
    ]
    ws.append(headers)

    # Style Header Row
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border
    ws.row_dimensions[1].height = 28

    # Populate Data
    row_idx = 2
    for u in users:
        role_str = str(u.role.value if hasattr(u.role, "value") else u.role)
        gender_str = str(u.gender.value if hasattr(u.gender, "value") else u.gender)
        full_name = f"{u.first_name} {u.last_name}".strip()

        row_data = [
            str(u.id),
            full_name or u.username,
            u.first_name,
            u.last_name,
            u.username,
            u.email,
            u.number or "",
            role_str,
            gender_str,
            "Yes" if u.is_verified else "No",
            u.streak or 0,
            u.university or "",
            u.last_login_date.strftime("%Y-%m-%d") if u.last_login_date else "",
            u.created_at.strftime("%Y-%m-%d %H:%M:%S") if u.created_at else "",
        ]
        ws.append(row_data)

        # Style data cells
        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = data_font
            c.border = cell_border
            if col_idx in (1, 8, 9, 10, 11, 13, 14):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[row_idx].height = 20
        row_idx += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@router.get("/users/export")
def export_platform_users(
    format: str = "xlsx",
    role: Optional[str] = None,
    current_admin: models.User = Depends(get_current_super_admin),
    db: Session = Depends(get_db),
):
    """
    Exports all current platform users to a formatted Excel workbook (.xlsx) or CSV file.
    """
    query = db.query(models.User).order_by(models.User.created_at.desc())
    if isinstance(role, str) and role.strip() and role.lower() != "all":
        r_clean = role.strip().lower()
        matched_enum = None
        for r_enum in Roles:
            if r_enum.value.lower() == r_clean or r_enum.name.lower() == r_clean:
                matched_enum = r_enum
                break
        if matched_enum:
            query = query.filter(
                or_(
                    models.User.role == matched_enum,
                    models.User.role == matched_enum.value,
                    func.lower(func.cast(models.User.role, String)) == r_clean,
                )
            )
        else:
            query = query.filter(func.lower(func.cast(models.User.role, String)) == r_clean)

    users = query.all()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    fmt = (format or "xlsx").strip().lower()

    if fmt == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "User ID", "Full Name", "First Name", "Last Name", "Username", "Email",
            "Phone Number", "Role", "Gender", "Verified", "Streak", "University",
            "Last Login Date", "Created At"
        ])
        for u in users:
            role_str = str(u.role.value if hasattr(u.role, "value") else u.role)
            gender_str = str(u.gender.value if hasattr(u.gender, "value") else u.gender)
            writer.writerow([
                str(u.id), f"{u.first_name} {u.last_name}".strip(), u.first_name, u.last_name,
                u.username, u.email, u.number or "", role_str, gender_str,
                "Yes" if u.is_verified else "No", u.streak or 0, u.university or "",
                u.last_login_date.strftime("%Y-%m-%d") if u.last_login_date else "",
                u.created_at.strftime("%Y-%m-%d %H:%M:%S") if u.created_at else "",
            ])
        output.seek(0)
        return StreamingResponse(
            io.BytesIO(output.getvalue().encode("utf-8")),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=nu_age_users_{timestamp}.csv"},
        )

    # Default: Excel (.xlsx)
    excel_stream = _generate_excel_workbook(users)
    return StreamingResponse(
        excel_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=nu_age_users_{timestamp}.xlsx"},
    )


@router.get("/users/{user_id}", response_model=AdminUserItem)
def get_user_details(
    user_id: UUID,
    current_admin: models.User = Depends(get_current_super_admin),
    db: Session = Depends(get_db),
):
    """
    Returns complete details of a single user by ID.
    """
    u = db.query(models.User).filter(models.User.id == user_id).first()
    if not u:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found.")

    role_val = str(u.role.value if hasattr(u.role, "value") else u.role)
    gender_val = str(u.gender.value if hasattr(u.gender, "value") else u.gender)
    full_name = f"{u.first_name} {u.last_name}".strip()

    return AdminUserItem(
        id=u.id,
        name=full_name or u.username,
        first_name=u.first_name,
        last_name=u.last_name,
        username=u.username,
        email=u.email,
        number=u.number,
        role=role_val,
        gender=gender_val,
        university=u.university,
        streak=u.streak or 0,
        is_verified=bool(u.is_verified),
        last_login_date=u.last_login_date.strftime("%Y-%m-%d") if u.last_login_date else None,
        created_at=u.created_at.strftime("%Y-%m-%d %H:%M") if u.created_at else None,
        enrolled_courses_count=len(u.courses) if hasattr(u, "courses") and u.courses else 0,
        created_courses_count=len(u.created_courses) if hasattr(u, "created_courses") and u.created_courses else 0,
        organisations_count=len(u.organisations) if hasattr(u, "organisations") and u.organisations else 0,
    )


@router.patch("/users/{user_id}")
def update_user_attributes(
    user_id: UUID,
    payload: AdminUserUpdateRequest,
    current_admin: models.User = Depends(get_current_super_admin),
    db: Session = Depends(get_db),
):
    """
    Modifies a user's role, verification status, streak, or university.
    """
    u = db.query(models.User).filter(models.User.id == user_id).first()
    if not u:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found.")

    if payload.role is not None:
        r_clean = payload.role.strip().capitalize()
        for r_enum in Roles:
            if r_enum.value.lower() == r_clean.lower():
                u.role = r_enum
                break

    if payload.is_verified is not None:
        u.is_verified = payload.is_verified

    if payload.streak is not None:
        u.streak = max(0, payload.streak)

    if payload.university is not None:
        u.university = payload.university.strip()

    db.commit()
    db.refresh(u)
    return {"message": "User updated successfully.", "user_id": str(u.id)}


@router.delete("/users/{user_id}")
def delete_user_account(
    user_id: UUID,
    current_admin: models.User = Depends(get_current_super_admin),
    db: Session = Depends(get_db),
):
    """
    Permanently deletes a user account with cascading cleanup of foreign-key dependencies
    (refresh tokens, device tokens, memberships, submissions, connections).
    """
    if str(current_admin.id) == str(user_id):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Cannot delete your own super administrator account while logged in.",
        )

    u = db.query(models.User).filter(models.User.id == user_id).first()
    if not u:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found.")

    try:
        # 1. Clean up refresh tokens (foreign key with no default cascade)
        db.query(models.RefreshToken).filter(models.RefreshToken.user_id == user_id).delete()

        # 2. Clean up device tokens
        db.query(models.DeviceToken).filter(models.DeviceToken.user_id == user_id).delete()

        # 3. Clean up OTP records if present
        if u.email:
            db.query(models.SignupOTP).filter(models.SignupOTP.email == u.email).delete()
            db.query(models.PasswordResetOTP).filter(models.PasswordResetOTP.email == u.email).delete()

        # 4. Delete the user (SQLAlchemy / Postgres cascades handle remaining child relations)
        db.delete(u)
        db.commit()
        return {"message": f"User account '{u.username}' ({u.email}) permanently deleted.", "user_id": str(user_id)}
    except Exception as ex:
        db.rollback()
        print(f"[platform_admin] Error deleting user {user_id}: {ex}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to delete user account: {ex}",
        )





# ─────────────────────────────────────────────────────────────────────────────
# 5. BULK COMMUNICATIONS (EMAIL & PUSH NOTIFICATIONS)
# ─────────────────────────────────────────────────────────────────────────────

def _filter_users_by_audience(db: Session, audience: str) -> List[models.User]:
    aud = (audience or "all").lower().strip()
    query = db.query(models.User)
    if aud == "students":
        query = query.filter(models.User.role == Roles.STUDENT)
    elif aud == "teachers":
        query = query.filter(models.User.role == Roles.TEACHER)
    elif aud == "admins":
        query = query.filter(models.User.role == Roles.ADMIN)
    elif aud == "unverified":
        query = query.filter(models.User.is_verified == False)
    return query.all()


def _dispatch_bulk_email_task(emails: List[str], subject: str, html_body: str, sender_name: str):
    """
    Background batch dispatcher for bulk emails using Resend with rate limit pacing.
    """
    resend.api_key = Settings().RESEND_API_KEY
    sender_email = f"{sender_name} <support@nu-age.name.ng>"

    batch_size = 50
    for i in range(0, len(emails), batch_size):
        chunk = emails[i : i + batch_size]
        for email in chunk:
            try:
                params: resend.Emails.SendParams = {
                    "from": sender_email,
                    "to": [email],
                    "subject": subject,
                    "html": html_body,
                }
                resend.Emails.send(params)
            except Exception as e:
                print(f"[platform_admin] Failed email to {email}: {e}")
        # Pace batches to respect provider rate limits
        time.sleep(0.5)


@router.post("/broadcast/email")
def broadcast_bulk_email(
    payload: BulkEmailRequest,
    background_tasks: BackgroundTasks,
    current_admin: models.User = Depends(get_current_super_admin),
    db: Session = Depends(get_db),
):
    """
    Sends a bulk email to all users matching the audience filter via Resend in the background.
    """
    if not payload.subject.strip():
        raise HTTPException(status_code=400, detail="Subject cannot be empty.")
    if not payload.body_html.strip():
        raise HTTPException(status_code=400, detail="Email body cannot be empty.")

    target_users = _filter_users_by_audience(db, payload.audience)
    emails = [u.email for u in target_users if u.email and "@" in u.email]

    if not emails:
        return {"message": "No users found matching audience criteria.", "recipient_count": 0}

    background_tasks.add_task(
        _dispatch_bulk_email_task,
        emails,
        payload.subject.strip(),
        payload.body_html.strip(),
        payload.sender_name or "Tobi from Nu Age",
    )

    return {
        "message": f"Bulk email queued successfully for {len(emails)} recipients in audience '{payload.audience}'.",
        "recipient_count": len(emails),
        "audience": payload.audience,
    }


@router.post("/broadcast/push")
def broadcast_bulk_push_notification(
    payload: BulkPushRequest,
    current_admin: models.User = Depends(get_current_super_admin),
    db: Session = Depends(get_db),
):
    """
    Dispatches a rich push notification via OneSignal & Firebase FCM to targeted devices,
    recording an audit entry in the notification_broadcasts database table.
    """
    if not payload.title.strip() or not payload.body.strip():
        raise HTTPException(status_code=400, detail="Notification title and body are required.")

    # ── 1. Resolve Target Audience ───────────────────────────────────────────
    is_test_to_me = payload.audience == "test_me"
    if is_test_to_me:
        target_users = [current_admin]
        user_ids = [current_admin.id]
    elif payload.audience == "all":
        target_users = []
        user_ids = []
    else:
        target_users = _filter_users_by_audience(db, payload.audience)
        user_ids = [u.id for u in target_users]

    if not user_ids and not is_test_to_me and payload.audience != "all":
        return {"message": "No users found matching audience criteria.", "device_count": 0}

    # ── 1.5 In-App UserNotification Persistence (PostgreSQL) ──────────────────
    target_route_path = (payload.action_route or "/notifications").strip()
    if not target_route_path.startswith("/"):
        target_route_path = "/" + target_route_path

    try:
        from services.notifications import dispatch_notification
        if is_test_to_me:
            dispatch_notification(
                db=db,
                recipient_user_ids=[current_admin.id],
                title=payload.title.strip(),
                body=payload.body.strip(),
                category="announcement",
                action_route=target_route_path,
                sender_id=current_admin.id,
                data_payload={"route": target_route_path, "image_url": payload.image_url},
                send_push=False,
                allow_self_notify=True,
            )
        elif payload.audience == "all":
            all_u_rows = db.query(models.User.id).all()
            all_uids = [r[0] for r in all_u_rows if r[0]]
            for c_idx in range(0, len(all_uids), 500):
                dispatch_notification(
                    db=db,
                    recipient_user_ids=all_uids[c_idx : c_idx + 500],
                    title=payload.title.strip(),
                    body=payload.body.strip(),
                    category="announcement",
                    action_route=target_route_path,
                    sender_id=current_admin.id,
                    data_payload={"route": target_route_path, "image_url": payload.image_url},
                    send_push=False,
                    allow_self_notify=True,
                )
        elif user_ids:
            dispatch_notification(
                db=db,
                recipient_user_ids=user_ids,
                title=payload.title.strip(),
                body=payload.body.strip(),
                category="announcement",
                action_route=target_route_path,
                sender_id=current_admin.id,
                data_payload={"route": target_route_path, "image_url": payload.image_url},
                send_push=False,
                allow_self_notify=True,
            )
    except Exception as in_app_ex:
        print(f"[platform_admin] Warning creating in-app UserNotification records: {in_app_ex}")

    # ── 2. OneSignal Rich Push Dispatch ──────────────────────────────────────
    onesignal_sent = False
    onesignal_id = None
    onesignal_recipients = 0
    try:
        settings = Settings()
        app_id = settings.get_onesignal_app_id()
        api_key = settings.get_onesignal_rest_api_key()

        if not app_id or not api_key:
            print(f"[platform_admin] WARNING: OneSignal broadcast skipped. Missing credentials. "
                  f"ONESIGNAL_APP_ID={'configured (' + settings.mask_onesignal_app_id() + ')' if app_id else 'MISSING'}, "
                  f"ONESIGNAL_REST_API_KEY={'configured (' + settings.mask_onesignal_key() + ')' if api_key else 'MISSING'}")
        else:
            import httpx
            onesignal_url = "https://api.onesignal.com/notifications"
            headers = {
                "Authorization": f"Key {api_key}",
                "Content-Type": "application/json; charset=utf-8",
            }
            target_route_path = (payload.action_route or "/notifications").strip()
            if not target_route_path.startswith("/"):
                target_route_path = "/" + target_route_path

            onesignal_payload = {
                "app_id": app_id,
                "headings": {"en": payload.title.strip()},
                "contents": {"en": payload.body.strip()},
                "data": {"route": target_route_path},
                "priority": payload.priority,
                "ttl": payload.ttl_seconds,
            }

            if payload.subtitle and payload.subtitle.strip():
                onesignal_payload["subtitle"] = {"en": payload.subtitle.strip()}

            if payload.image_url and payload.image_url.strip():
                img = payload.image_url.strip()
                onesignal_payload["big_picture"] = img
                onesignal_payload["ios_attachments"] = {"banner": img}

            if payload.action_button_label and payload.action_button_label.strip():
                onesignal_payload["buttons"] = [
                    {"id": "btn_action", "text": payload.action_button_label.strip()}
                ]

            targeted_uids_str = []
            if is_test_to_me:
                targeted_uids_str = [str(current_admin.id)]
                onesignal_payload["include_aliases"] = {"external_id": targeted_uids_str}
                onesignal_payload["target_channel"] = "push"
            elif payload.audience == "all":
                onesignal_payload["included_segments"] = ["Subscribed Users"]
            else:
                targeted_uids_str = [str(uid) for uid in user_ids]
                onesignal_payload["include_aliases"] = {"external_id": targeted_uids_str}
                onesignal_payload["target_channel"] = "push"

            target_desc = f"segment 'Subscribed Users'" if payload.audience == "all" and not is_test_to_me else f"{len(targeted_uids_str)} aliases"
            print(f"[platform_admin] OneSignal broadcasting to {target_desc}... "
                  f"app_id={settings.mask_onesignal_app_id()}, auth=Key {settings.mask_onesignal_key()}")

            with httpx.Client(timeout=15.0) as client:
                res = client.post(onesignal_url, json=onesignal_payload, headers=headers)
                res_text = res.text
                try:
                    res_json = res.json()
                except Exception:
                    res_json = {}

                onesignal_id = res_json.get("id")
                onesignal_recipients = res_json.get("recipients", 0)
                os_errors = res_json.get("errors")
                os_warnings = res_json.get("warnings")

                if res.status_code in (200, 201) and not os_errors and onesignal_recipients > 0:
                    onesignal_sent = True
                    print(f"[platform_admin] OneSignal broadcast SUCCESS: id={onesignal_id}, recipients={onesignal_recipients}, warnings={os_warnings}")
                else:
                    print(f"[platform_admin] OneSignal broadcast issue: status={res.status_code}, id={onesignal_id}, recipients={onesignal_recipients}, errors={os_errors}, warnings={os_warnings}, raw={res_text}")

                    # 1. Fallback for broadcast "all" if "Subscribed Users" had 0 subscribers
                    if payload.audience == "all" and not is_test_to_me and (os_errors or onesignal_recipients == 0):
                        for alt_segment in [["Active Users"], ["Total Subscriptions"], ["All"]]:
                            print(f"[platform_admin] Retrying broadcast with alternative segment: {alt_segment}...")
                            fb_seg_payload = dict(onesignal_payload)
                            fb_seg_payload["included_segments"] = alt_segment
                            fb_seg_payload["target_channel"] = "push"
                            res_fb_seg = client.post(onesignal_url, json=fb_seg_payload, headers=headers)
                            try:
                                fb_seg_json = res_fb_seg.json()
                            except Exception:
                                fb_seg_json = {}
                            fb_seg_errors = fb_seg_json.get("errors")
                            fb_seg_recipients = fb_seg_json.get("recipients", 0)
                            if res_fb_seg.status_code in (200, 201) and not fb_seg_errors and fb_seg_recipients > 0:
                                onesignal_sent = True
                                onesignal_id = fb_seg_json.get("id") or onesignal_id
                                onesignal_recipients = fb_seg_recipients
                                print(f"[platform_admin] OneSignal broadcast SUCCESS via segment {alt_segment}: id={onesignal_id}, recipients={onesignal_recipients}")
                                break
                            else:
                                print(f"[platform_admin] OneSignal segment {alt_segment} response: status={res_fb_seg.status_code}, recipients={fb_seg_recipients}, errors={fb_seg_errors}")

                    # 2. Fallback retry for targeted aliases with legacy include_external_user_ids
                    if targeted_uids_str and (res.status_code == 400 or os_errors or onesignal_recipients == 0):
                        print(f"[platform_admin] Retrying targeted broadcast with legacy include_external_user_ids...")
                        fb_payload = dict(onesignal_payload)
                        fb_payload.pop("include_aliases", None)
                        fb_payload.pop("target_channel", None)
                        fb_payload["include_external_user_ids"] = targeted_uids_str
                        res_fb = client.post(onesignal_url, json=fb_payload, headers=headers)
                        try:
                            fb_json = res_fb.json()
                        except Exception:
                            fb_json = {}
                        fb_recipients = fb_json.get("recipients", 0)
                        if res_fb.status_code in (200, 201) and not fb_json.get("errors") and fb_recipients > 0:
                            onesignal_sent = True
                            onesignal_id = fb_json.get("id") or onesignal_id
                            onesignal_recipients = fb_recipients
                            print(f"[platform_admin] OneSignal targeted fallback SUCCESS: id={onesignal_id}, recipients={onesignal_recipients}")
                        else:
                            print(f"[platform_admin] OneSignal fallback response: status={res_fb.status_code}, id={onesignal_id}, recipients={fb_recipients}, errors={fb_json.get('errors')}")

    except Exception as os_ex:
        print(f"[platform_admin] OneSignal broadcast exception: {os_ex!r}")

    # ── 3. Direct FCM Device Token Broadcast (Fallback/Secondary) ─────────────
    if payload.audience == "all":
        tokens = db.query(models.DeviceToken).all()
    elif user_ids:
        tokens = db.query(models.DeviceToken).filter(models.DeviceToken.user_id.in_(user_ids)).all()
    else:
        tokens = []
    token_strings = [t.token for t in tokens if t.token]

    sent_count = 0
    failure_count = 0
    dead_tokens = []

    if token_strings:
        try:
            import firebase_admin
            from firebase_admin import messaging

            data_payload = {"route": payload.action_route or "/notifications"}

            chunk_size = 500
            for i in range(0, len(token_strings), chunk_size):
                chunk = token_strings[i : i + chunk_size]
                msg = messaging.MulticastMessage(
                    notification=messaging.Notification(
                        title=payload.title.strip(),
                        body=payload.body.strip(),
                        image=payload.image_url.strip() if payload.image_url else None,
                    ),
                    data=data_payload,
                    tokens=chunk,
                )
                response = messaging.send_each_for_multicast(msg)
                sent_count += response.success_count
                failure_count += response.failure_count

                if response.failure_count > 0:
                    for idx, resp in enumerate(response.responses):
                        if not resp.success:
                            if getattr(resp.exception, "code", "") == "messaging/registration-token-not-registered":
                                dead_tokens.append(chunk[idx])

            if dead_tokens:
                db.query(models.DeviceToken).filter(models.DeviceToken.token.in_(dead_tokens)).delete(synchronize_session=False)
                db.commit()

        except Exception as ex:
            print(f"[platform_admin] Direct FCM broadcast warning: {ex}")

    # ── 4. Record Broadcast Audit Log ─────────────────────────────────────────
    total_targeted = onesignal_recipients or len(token_strings) or (len(user_ids) if user_ids else 0)
    if not total_targeted and payload.audience == "all":
        try:
            total_targeted = db.query(models.User).count()
        except Exception:
            total_targeted = 0
    delivery_status = "test" if is_test_to_me else ("completed" if (onesignal_sent or sent_count > 0) else "failed")

    try:
        broadcast_record = models.NotificationBroadcast(
            sender_id=current_admin.id,
            sender_username=current_admin.username,
            title=payload.title.strip(),
            body=payload.body.strip(),
            subtitle=payload.subtitle.strip() if payload.subtitle else None,
            image_url=payload.image_url.strip() if payload.image_url else None,
            action_route=payload.action_route,
            action_button_label=payload.action_button_label,
            audience=payload.audience,
            targeted_devices_count=total_targeted,
            delivered_count=sent_count or total_targeted if onesignal_sent else 0,
            failed_count=failure_count,
            priority=payload.priority,
            ttl_seconds=payload.ttl_seconds,
            onesignal_id=onesignal_id,
            status=delivery_status,
        )
        db.add(broadcast_record)
        db.commit()
    except Exception as db_err:
        print(f"[platform_admin] Failed logging broadcast to database: {db_err}")
        db.rollback()

    status_msg = (
        "Test notification dispatched to your device."
        if is_test_to_me
        else f"Broadcast sent! Targeted ~{total_targeted:,} recipients (OneSignal: {'OK' if onesignal_sent else 'Standby'}, Direct FCM: {sent_count})."
    )

    return {
        "message": status_msg,
        "devices_targeted": total_targeted,
        "sent_count": sent_count,
        "onesignal_sent": onesignal_sent,
        "onesignal_id": onesignal_id,
        "status": delivery_status,
    }


@router.get("/broadcast/history", response_model=BroadcastHistoryResponse)
def get_broadcast_history(
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=50),
    current_admin: models.User = Depends(get_current_super_admin),
    db: Session = Depends(get_db),
):
    """
    Returns a paginated audit log of all push notification broadcasts dispatched
    from the Super Admin Center.
    """
    query = db.query(models.NotificationBroadcast).order_by(desc(models.NotificationBroadcast.created_at))
    total = query.count()
    records = query.offset((page - 1) * limit).limit(limit).all()

    items = []
    for r in records:
        items.append(
            BroadcastHistoryItem(
                id=r.id,
                sender_username=r.sender_username,
                title=r.title,
                body=r.body,
                subtitle=r.subtitle,
                image_url=r.image_url,
                action_route=r.action_route,
                action_button_label=r.action_button_label,
                audience=r.audience,
                targeted_devices_count=r.targeted_devices_count or 0,
                delivered_count=r.delivered_count or 0,
                failed_count=r.failed_count or 0,
                priority=r.priority or 10,
                ttl_seconds=r.ttl_seconds or 86400,
                status=r.status or "completed",
                created_at=r.created_at.isoformat() if r.created_at else None,
            )
        )

    return BroadcastHistoryResponse(
        items=items,
        total=total,
        page=page,
        limit=limit,
    )



# ─────────────────────────────────────────────────────────────────────────────
# 6. SYSTEM HEALTH & DIAGNOSTICS
# ─────────────────────────────────────────────────────────────────────────────

@router.get("/health")
def platform_system_health(
    current_admin: models.User = Depends(get_current_super_admin),
    db: Session = Depends(get_db),
):
    """
    Checks database connection responsiveness, server timestamp, and active pool health.
    """
    t0 = time.time()
    try:
        from sqlalchemy import text
        db.execute(text("SELECT 1;"))
        db_latency_ms = round((time.time() - t0) * 1000, 2)
        db_status = "healthy"
    except Exception as ex:
        db_latency_ms = round((time.time() - t0) * 1000, 2)
        db_status = f"unhealthy: {ex}"

    return {
        "status": "online",
        "server_time_utc": datetime.now(timezone.utc).isoformat(),
        "database_status": db_status,
        "database_latency_ms": db_latency_ms,
        "super_admin": current_admin.username,
    }


@router.get("/onesignal/diagnostics")
def onesignal_diagnostics(
    current_admin: models.User = Depends(get_current_super_admin),
):
    """
    Directly queries OneSignal REST API to inspect:
    1. App configuration and active subscriber counts
    2. Exact segment names configured in OneSignal
    3. Last notifications sent (including successful dashboard messages)
    """
    import httpx
    settings = Settings()
    app_id = settings.get_onesignal_app_id()
    api_key = settings.get_onesignal_rest_api_key()

    if not app_id or not api_key:
        return {
            "error": "Missing credentials in Settings",
            "app_id_configured": bool(app_id),
            "api_key_configured": bool(api_key),
        }

    headers = {
        "Authorization": f"Key {api_key}",
        "Content-Type": "application/json; charset=utf-8",
    }

    result = {
        "app_id": settings.mask_onesignal_app_id(),
        "key": settings.mask_onesignal_key(),
    }

    with httpx.Client(timeout=15.0) as client:
        # 1. Fetch App Info & Subscribers
        try:
            r_app = client.get(f"https://api.onesignal.com/apps/{app_id}", headers=headers)
            result["app_info"] = r_app.json() if r_app.status_code == 200 else {"status": r_app.status_code, "text": r_app.text}
        except Exception as e:
            result["app_info"] = {"error": str(e)}

        # 2. Fetch Segments
        try:
            r_seg = client.get(f"https://api.onesignal.com/apps/{app_id}/segments", headers=headers)
            result["segments"] = r_seg.json() if r_seg.status_code == 200 else {"status": r_seg.status_code, "text": r_seg.text}
        except Exception as e:
            result["segments"] = {"error": str(e)}

        # 3. Fetch Recent Notifications
        try:
            r_notifs = client.get(f"https://api.onesignal.com/notifications?app_id={app_id}&limit=5", headers=headers)
            result["recent_notifications"] = r_notifs.json() if r_notifs.status_code == 200 else {"status": r_notifs.status_code, "text": r_notifs.text}
        except Exception as e:
            result["recent_notifications"] = {"error": str(e)}

    print(f"[platform_admin] OneSignal diagnostics fetched: app_name={result.get('app_info', {}).get('name') if isinstance(result.get('app_info'), dict) else 'N/A'}")
    return result
